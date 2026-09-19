"""
Excel processing (openpyxl).

Reading : every worksheet is scanned for natural-language string cells.
Writing : the workbook is COPIED, then only the identified text cells are
          replaced in the copy.  Sheet names, sheet order, row/column order,
          numbers, dates, booleans, formulas and cell formatting all survive,
          and the original workbook is never written to.
"""

from __future__ import annotations

import re
import shutil
from datetime import date, datetime, time
from pathlib import Path

from config import OUTPUT_DIR
from core.errors import CorruptFileError, EmptyContentError, ExportError, ExtractionError
from processors import (
    ExtractedDocument,
    normalize_text,
    resolve_input,
    unique_output_path,
)

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - dependency guard
    raise ExtractionError(
        "openpyxl is not installed.", "Run: pip install -r requirements.txt"
    ) from exc

MIN_WORDS_PER_CELL = 4
NUMERIC_LIKE = re.compile(r"^[\s$€£₹%+-]*[\d.,]+\s*[%a-zA-Z]{0,3}$")
DATE_LIKE = re.compile(r"^\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}")


def _open(source: Path, read_only: bool = False):
    try:
        return load_workbook(source, data_only=False, read_only=read_only)
    except Exception as exc:
        raise CorruptFileError(
            f"'{source.name}' could not be opened as an Excel workbook.",
            "It may be corrupted, password protected, or an old .xls file. "
            "Save it as .xlsx and try again.",
        ) from exc


def is_rewritable_value(value) -> bool:
    """True only for string cells that read like prose."""
    if isinstance(value, (int, float, bool, datetime, date, time)) or value is None:
        return False
    text = str(value).strip()
    if not text or text.startswith("="):            # never touch a formula
        return False
    if NUMERIC_LIKE.match(text) or DATE_LIKE.match(text):
        return False
    if text.startswith(("http://", "https://", "www.")):
        return False
    return len(text.split()) >= MIN_WORDS_PER_CELL


def extract(path: str | Path) -> ExtractedDocument:
    """Collect every natural-language cell across every worksheet."""
    source = resolve_input(path)
    workbook = _open(source)

    segments: list[dict] = []
    sheet_summary: list[dict] = []
    formula_count = 0

    try:
        for sheet in workbook.worksheets:
            sheet_cells = 0
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        formula_count += 1
                        continue
                    if is_rewritable_value(value):
                        text = str(value).strip()
                        segments.append(
                            {
                                "kind": "cell",
                                "sheet": sheet.title,
                                "coordinate": cell.coordinate,
                                "row": cell.row,
                                "column": cell.column,
                                "text": text,
                                "word_count": len(text.split()),
                            }
                        )
                        sheet_cells += 1
            sheet_summary.append(
                {
                    "sheet": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "rewritable_cells": sheet_cells,
                }
            )
    finally:
        workbook.close()

    if not sheet_summary:
        raise EmptyContentError(f"'{source.name}' contains no worksheets.")

    combined = normalize_text("\n\n".join(s["text"] for s in segments))
    warnings: list[str] = []
    if not segments:
        warnings.append(
            "No natural-language cells were found; this workbook looks purely "
            "numeric or made of short labels."
        )

    return ExtractedDocument(
        text=combined,
        source_path=source,
        file_type="xlsx",
        segments=segments,
        metadata={
            "file_name": source.name,
            "sheet_count": len(sheet_summary),
            "sheets": sheet_summary,
            "formula_cells": formula_count,
            "rewritable_cells": len(segments),
        },
        warnings=warnings,
    )


def write_rewritten(
    source_path: str | Path,
    replacements: dict[tuple[str, str], str],
    stem: str | None = None,
    output_dir: Path | None = None,
) -> tuple[Path, list[dict]]:
    """
    Write a NEW workbook with the given cell replacements applied.

    `replacements` is keyed by `(sheet_name, coordinate)`, e.g. ("Sheet1", "B4").
    Returns `(path, applied_changes)`.
    """
    source = resolve_input(source_path)
    directory = Path(output_dir) if output_dir else OUTPUT_DIR
    stem = stem or f"{source.stem}_natural_rewrite"
    destination = unique_output_path(directory, stem, source.suffix.lower())

    try:
        shutil.copyfile(source, destination)         # original stays untouched
    except OSError as exc:
        raise ExportError("Could not create the output workbook.", str(exc)) from exc

    workbook = _open(destination)
    applied: list[dict] = []

    try:
        for (sheet_name, coordinate), new_value in replacements.items():
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            cell = sheet[coordinate]
            old_value = cell.value
            if not is_rewritable_value(old_value):
                continue
            new_value = (new_value or "").strip()
            if not new_value or new_value == str(old_value).strip():
                continue
            cell.value = new_value
            applied.append(
                {
                    "sheet": sheet_name,
                    "cell": coordinate,
                    "original": str(old_value),
                    "rewritten": new_value,
                }
            )
        workbook.save(destination)
    except OSError as exc:
        raise ExportError("Could not save the output workbook.", str(exc)) from exc
    finally:
        workbook.close()

    return destination, applied
