"""
Create sample input files for testing.

Run:  python make_samples.py

Writes one file per supported type into `input/` so every menu option in
`app.py` can be exercised immediately, with or without an API key.
"""

from __future__ import annotations

from pathlib import Path

from config import INPUT_DIR, ensure_directories

SAMPLE_TEXT = """The Role of Automation in Small Business Operations

In today's fast-paced world, it is important to note that automation has become
an essential component of modern business operations. Furthermore, organizations
that leverage automation tools are able to utilize their resources in order to
achieve greater operational efficiency. Moreover, the majority of small
businesses report significant improvements after adopting automated workflows.

Automation can be applied to a wide range of business processes. Additionally,
invoicing, customer follow-up, inventory tracking and reporting are commonly
automated. It should be noted that a company in Lahore reduced its manual data
entry by 42% between January 2024 and March 2025 after deploying a simple
workflow system.

Furthermore, it is important to note that automation is not a replacement for
human judgement. Moreover, the most effective implementations combine automated
processes with human oversight. In conclusion, businesses that approach
automation carefully are able to achieve meaningful and sustainable results.
"""

SAMPLE_CSV_ROWS = [
    ["id", "customer", "rating", "signup_date", "feedback"],
    [
        "1001", "Ayesha Khan", "4", "2025-03-11",
        "It is important to note that the onboarding process was quite smooth and "
        "the support team responded in a timely manner.",
    ],
    [
        "1002", "Bilal Ahmed", "2", "2025-04-02",
        "Furthermore, the mobile application was slow to load and additionally the "
        "checkout process failed on two separate occasions.",
    ],
    [
        "1003", "Sana Malik", "5", "2025-04-19",
        "The product exceeded my expectations in terms of reliability and the "
        "documentation was comprehensive and easy to follow.",
    ],
    ["1004", "Usman Tariq", "3", "2025-05-07", "Average"],
]


def create_samples() -> list[str]:
    """Create one sample file per supported type. Returns the paths created."""
    ensure_directories()
    created: list[str] = []

    # -- TXT ---------------------------------------------------------------
    txt_path = INPUT_DIR / "sample.txt"
    txt_path.write_text(SAMPLE_TEXT, encoding="utf-8")
    created.append(str(txt_path))

    # -- CSV ---------------------------------------------------------------
    import csv

    csv_path = INPUT_DIR / "sample.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerows(SAMPLE_CSV_ROWS)
    created.append(str(csv_path))

    # -- DOCX --------------------------------------------------------------
    try:
        import docx

        document = docx.Document()
        document.add_heading("The Role of Automation in Small Business Operations", 1)
        for paragraph in SAMPLE_TEXT.split("\n\n")[1:]:
            document.add_paragraph(paragraph.strip().replace("\n", " "))
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Process"
        table.cell(0, 1).text = "Notes"
        table.cell(1, 0).text = "Invoicing"
        table.cell(1, 1).text = (
            "Furthermore, invoicing is one of the most commonly automated processes "
            "in small organizations."
        )
        docx_path = INPUT_DIR / "sample.docx"
        document.save(str(docx_path))
        created.append(str(docx_path))
    except ImportError:
        print("python-docx not installed; skipped sample.docx")

    # -- XLSX --------------------------------------------------------------
    try:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Feedback"
        for row in SAMPLE_CSV_ROWS:
            sheet.append(row)
        sheet["F1"] = "score_x2"
        sheet["F2"] = "=C2*2"
        summary = workbook.create_sheet("Summary")
        summary["A1"] = "Observation"
        summary["A2"] = (
            "It is important to note that overall satisfaction improved during the "
            "period under review."
        )
        summary["B1"] = "Total"
        summary["B2"] = "=SUM(Feedback!C2:C5)"
        xlsx_path = INPUT_DIR / "sample.xlsx"
        workbook.save(str(xlsx_path))
        created.append(str(xlsx_path))
    except ImportError:
        print("openpyxl not installed; skipped sample.xlsx")

    # -- PDF ---------------------------------------------------------------
    try:
        from processors.pdf_processor import write_pdf

        pdf_path = write_pdf(
            SAMPLE_TEXT,
            "sample",
            output_dir=INPUT_DIR,
            title="Automation in Small Business Operations",
        )
        created.append(str(pdf_path))
    except Exception as exc:
        print(f"Could not create sample.pdf: {exc}")

    return created


if __name__ == "__main__":
    for item in create_samples():
        print("created:", item)
    print(f"\nSample files are in: {Path(INPUT_DIR)}")
